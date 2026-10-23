"""导出待上架的套图：每套一个文件夹（01.jpg 是首图），外加一张汇总表，方便上架 Wildberries。"""
from __future__ import annotations

import re
import shutil
import time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import store


def _safe(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|\r\n]+', "_", str(name)).strip()[:40] or "商品"


def export_sets(items: list[dict], job) -> None:
    """items: [{pid, sid}]"""
    root = store.EXPORTS / time.strftime("%Y%m%d-%H%M%S")
    root.mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "待上架"
    head = ["文件夹", "1688商品ID", "中文标题", "俄语标题", "颜色/款式", "尺码", "价格(元)", "库存", "Logo",
            "卖点(俄)", "卖点(中)", "图片数", "1688链接"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4F46E5")
    job.total = len(items)
    for it in items:
        p = store.get_product(it["pid"])
        st = next((x for x in (p or {}).get("sets", []) if x["id"] == it["sid"]), None)
        if not st:
            job.done += 1
            continue
        pdir = store.product_dir(p["id"])
        folder = root / _safe(f"{p.get('offer_id') or p['id']}_{st['sku']}_{st.get('logo_name') or '无标'}")
        folder.mkdir(exist_ok=True)
        n = 0
        for im in st["images"]:
            if im.get("file") and (pdir / im["file"]).exists():
                n += 1
                shutil.copy(pdir / im["file"], folder / f"{n:02d}.jpg")
        info = p.get("ai") or {}
        sizes = p.get("sizes") or [""]
        for size in sizes:
            sku = next((k for k in p.get("skus") or []
                        if st["sku"] in (k.get("parts") or []) and (not size or size in (k.get("parts") or []))), {})
            ws.append([folder.name, p.get("offer_id") or "", p.get("title", ""), info.get("title_ru", ""), st["sku"],
                       size, sku.get("price", p.get("price_min")), sku.get("stock", ""), st.get("logo_name", ""),
                       " / ".join(info.get("selling_points_ru") or []), " / ".join(info.get("selling_points_cn") or []),
                       n, p.get("url", "")])
        job.done += 1
    for i, w in enumerate([36, 16, 40, 40, 14, 8, 10, 8, 12, 40, 30, 8, 46]):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="center")
    wb.save(root / "上架汇总.xlsx")
    job.results.append(str(root))
    job.finish(f"导出完成：{root}")
